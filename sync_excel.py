"""
SEPTEMBER 2026 MASTER EXCEL SYNC SCRIPT FOR AWARD CHOICES & RECOMMENDATIONS
===========================================================================
Run this script to update 'Exium_Award_Choice_Master_Sep_2026.xlsx' and
'Calculation\\September_2026_Award_Calculation_Master.xlsx' from a saved JSON file.
"""
import openpyxl
import json
import os
import sys

BASE_DIR = r"G:\Exium\2026\Award\September"
EXCEL_FILES = [
    os.path.join(BASE_DIR, "Exium_Award_Choice_Master_Sep_2026.xlsx"),
    os.path.join(BASE_DIR, "Calculation", "September_2026_Award_Calculation_Master.xlsx"),
]
DEFAULT_CHOICES_FILE = os.path.join(BASE_DIR, "saved_award_choices_sep_2026.json")


def sync_workbook(excel_path, choices_data):
    if not os.path.exists(excel_path):
        print(f"Skipping (not found): {excel_path}")
        return

    wb = openpyxl.load_workbook(excel_path)

    # Update Award_Exium_Sep_2026 & Award_Coralcal_Sep_2026
    # Columns: Col 8 = SAP Area Code, Col 10 = SAP MIO Code,
    # Col 23 = RH Recommendation, Col 24 = RH Cancel Reason,
    # Col 25 = ZH Recommendation, Col 26 = ZH Cancel Reason,
    # Col 27 = Selected Gift Voucher, Col 28 = Status, Col 29 = Timestamp (BDT)
    for sheet_name, prefix in [("Award_Exium_Sep_2026", "EX_"), ("Award_Coralcal_Sep_2026", "CC_")]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        updated = 0
        for r in range(2, ws.max_row + 1):
            tty = str(ws.cell(r, 8).value or "").strip()
            mio = str(ws.cell(r, 10).value or "").strip()
            rec_id = f"{prefix}{tty}_{mio}"
            item = choices_data.get(rec_id)
            if not item:
                continue

            rh_rec = item.get("rh_rec", "Recommended")
            rh_reason = item.get("rh_reason", "")
            zh_rec = item.get("zh_rec", "Recommended")
            zh_reason = item.get("zh_reason", "")
            voucher = item.get("voucher", "")
            status = item.get("status", "Complete" if voucher else "Pending")
            ts = item.get("timestamp", "")

            ws.cell(r, 23, value=rh_rec)
            ws.cell(r, 24, value=rh_reason)
            ws.cell(r, 25, value=zh_rec)
            ws.cell(r, 26, value=zh_reason)
            ws.cell(r, 27, value="" if status == "Cancelled" else voucher)
            ws.cell(r, 28, value=status)
            ws.cell(r, 29, value=ts)
            updated += 1
        print(f"[{os.path.basename(excel_path)}] Updated {updated} rows in {sheet_name}")

    wb.save(excel_path)
    print(f"Saved: {excel_path}")


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_CHOICES_FILE
    if os.path.exists(src):
        with open(src, "r", encoding="utf-8") as f:
            data = json.load(f)
        for p in EXCEL_FILES:
            sync_workbook(p, data)
    else:
        print(f"Choices file '{src}' not found. Provide a JSON file path to sync.")
