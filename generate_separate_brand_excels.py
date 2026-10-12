"""
GENERATE SEPARATE EXCEL CALCULATION WORKBOOKS FOR EXIUM AND CORALCAL (SEPTEMBER 2026)
=====================================================================================
Creates:
1. G:\\Exium\\2026\\Award\\September\\Calculation\\Exium_Award_Calculation_Sep_2026.xlsx
   (and copy in G:\\Exium\\2026\\Award\\September\\Exium_Award_Calculation_Sep_2026.xlsx)
2. G:\\Exium\\2026\\Award\\September\\Calculation\\Coralcal_Award_Calculation_Sep_2026.xlsx
   (and copy in G:\\Exium\\2026\\Award\\September\\Coralcal_Award_Calculation_Sep_2026.xlsx)
"""

import os
import json
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = r"G:\Exium\2026\Award\September"
CALC_DIR = os.path.join(BASE_DIR, "Calculation")
DATA_JSON = os.path.join(BASE_DIR, "september_award_data.json")


def build_single_brand_workbook(records, brand_title, prefix, primary_hex, soft_hex, out_path):
    wb = openpyxl.Workbook()

    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_sec = Font(name="Calibri", size=11, bold=True, color="1E293B")
    font_hdr = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, bold=False, color="0F172A")
    font_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
    font_green = Font(name="Calibri", size=10, bold=True, color="065F46")
    font_blue = Font(name="Calibri", size=10, bold=True, color="1E40AF")
    font_red = Font(name="Calibri", size=10, bold=True, color="991B1B")

    fill_navy = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_brand = PatternFill(start_color=primary_hex, end_color=primary_hex, fill_type="solid")
    fill_soft = PatternFill(start_color=soft_hex, end_color=soft_hex, fill_type="solid")
    fill_green = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    fill_blue_soft = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    fill_red_soft = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    fill_yellow = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    thin = Side(border_style="thin", color="CBD5E1")
    border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

    align_c = Alignment(horizontal="center", vertical="center")
    align_cw = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_l = Alignment(horizontal="left", vertical="center")
    align_r = Alignment(horizontal="right", vertical="center")

    # Stats
    total_ttys = len(records)
    c1_pass = sum(1 for r in records if r[f"{prefix}_c1_rx"])
    c2b_pass = sum(1 for r in records if r["tot_sales_positive"])
    c2c_pass = sum(1 for r in records if r[f"{prefix}_growth_met"])
    tier_3000 = sum(1 for r in records if r[f"{prefix}_award"] == 3000)
    tier_1000 = sum(1 for r in records if r[f"{prefix}_award"] == 1000)
    total_ach = tier_3000 + tier_1000
    no_award = total_ttys - total_ach
    val_3000 = tier_3000 * 3000
    val_1000 = tier_1000 * 1000
    total_val = val_3000 + val_1000

    # =========================================================================
    # SHEET 1: Summary_Dashboard
    # =========================================================================
    ws_sum = wb.active
    ws_sum.title = f"{prefix.upper()}_Summary_Dashboard"
    ws_sum.views.sheetView[0].showGridLines = True

    ws_sum.merge_cells("B2:J2")
    c = ws_sum["B2"]
    c.value = f"{brand_title.upper()} — SEPTEMBER 2026 AWARD CALCULATION SUMMARY"
    c.font = font_title
    c.fill = fill_brand
    c.alignment = align_c
    ws_sum.row_dimensions[2].height = 34

    kpi_headers = ["Metric / Criteria", "Territory / Achiever Count", "Award Value", "Remarks"]
    ws_sum.row_dimensions[4].height = 26
    for i, h in enumerate(kpi_headers, start=2):
        cell = ws_sum.cell(4, i, h)
        cell.font = font_hdr
        cell.fill = fill_navy
        cell.alignment = align_c
        cell.border = border_all

    kpi_rows = [
        ("Total Territories Evaluated (FF List)", total_ttys, "-", "Across 35 Zones & 252 Regions"),
        ("Criteria 1 Met: Seen Rx >= 200 (Full Month Sep'26)", c1_pass, "-", "Eligible for at least 1,000 Award"),
        ("Criteria 2(b) Met: Total Territory Sales Positive (Sep'26 > Aug'26)", c2b_pass, "-", "From Award Data, RPL.xlsx"),
        (f"Criteria 2(c) Met: {brand_title} Value Growth Met (30-Box Target)", c2c_pass, "-", "From Value Growth Calculation - Team A.xlsx"),
        ("3,000 Award Achievers (All 3 Criteria Met: 1 + 2b + 2c)", tier_3000, val_3000, "3,000 Gift Voucher"),
        ("1,000 Award Achievers (Criteria 1 Only: Seen Rx >= 200)", tier_1000, val_1000, "1,000 Gift Voucher"),
        (f"Total Eligible {brand_title} Award Achievers", total_ach, total_val, "1,000 + 3,000 Combined"),
        ("Not Eligible (Seen Rx < 200)", no_award, 0, "No Award"),
    ]

    for idx, (m_label, cnt, val, rem) in enumerate(kpi_rows, start=5):
        ws_sum.row_dimensions[idx].height = 22
        vals = [m_label, cnt, val, rem]
        for c_i, v in enumerate(vals, start=2):
            cell = ws_sum.cell(idx, c_i, v)
            cell.border = border_all
            cell.font = font_bold if idx in (9, 10, 11) else font_data
            if idx == 11:
                cell.fill = fill_yellow
            elif idx == 9:
                cell.fill = fill_green
            elif idx == 10:
                cell.fill = fill_blue_soft
            cell.alignment = align_l if c_i in (2, 5) else align_c
            if isinstance(v, int):
                cell.number_format = "#,##0"

    # Zone-wise breakdown table
    r_start = 15
    ws_sum.cell(r_start - 1, 2, f"Zone-Wise {brand_title} Award Breakdown (35 Zones)").font = font_sec
    z_headers = [
        "SL", "Zone Name", "SAP Zone Code", "Zonal Head",
        "Total Territories", "Seen Rx >= 200", "3,000 Achievers",
        "1,000 Achievers", "Total Achievers", "No Award (<200 Rx)", "Total Award Value"
    ]
    ws_sum.row_dimensions[r_start].height = 28
    for c_i, h in enumerate(z_headers, start=1):
        cell = ws_sum.cell(r_start, c_i, h)
        cell.font = font_hdr
        cell.fill = fill_brand
        cell.alignment = align_cw
        cell.border = border_all

    zones_dict = {}
    for r in records:
        z = r["zone_name"]
        if z not in zones_dict:
            zones_dict[z] = {
                "zone_name": z,
                "sap_zone": r["sap_zone_code"],
                "zh": r["zm_name"],
                "ttys": 0,
                "rx200": 0,
                "a3000": 0,
                "a1000": 0,
                "no_aw": 0,
                "val": 0
            }
        zd = zones_dict[z]
        zd["ttys"] += 1
        if r[f"{prefix}_c1_rx"]:
            zd["rx200"] += 1
        if r[f"{prefix}_award"] == 3000:
            zd["a3000"] += 1
            zd["val"] += 3000
        elif r[f"{prefix}_award"] == 1000:
            zd["a1000"] += 1
            zd["val"] += 1000
        else:
            zd["no_aw"] += 1

    r_curr = r_start + 1
    for sl, zd in enumerate(sorted(zones_dict.values(), key=lambda x: x["zone_name"]), start=1):
        ws_sum.row_dimensions[r_curr].height = 20
        tot_z_ach = zd["a3000"] + zd["a1000"]
        row_v = [
            sl, zd["zone_name"], zd["sap_zone"], zd["zh"],
            zd["ttys"], zd["rx200"], zd["a3000"], zd["a1000"],
            tot_z_ach, zd["no_aw"], zd["val"]
        ]
        for c_i, v in enumerate(row_v, start=1):
            cell = ws_sum.cell(r_curr, c_i, v)
            cell.border = border_all
            cell.font = font_bold if c_i in (9, 11) else font_data
            if sl % 2 == 0:
                cell.fill = fill_zebra
            cell.alignment = align_l if c_i in (2, 4) else align_c
            if c_i >= 5 and isinstance(v, int):
                cell.number_format = "#,##0"
        r_curr += 1

    ws_sum.row_dimensions[r_curr].height = 24
    tot_row = [
        "", "NATIONAL GRAND TOTAL", "", "",
        total_ttys, c1_pass, tier_3000, tier_1000, total_ach, no_award, total_val
    ]
    for c_i, v in enumerate(tot_row, start=1):
        cell = ws_sum.cell(r_curr, c_i, v)
        cell.border = border_all
        cell.font = font_bold
        cell.fill = fill_yellow
        cell.alignment = align_l if c_i == 2 else align_c
        if c_i >= 5 and isinstance(v, int):
            cell.number_format = "#,##0"

    sum_widths = [6, 34, 16, 26, 16, 16, 16, 16, 16, 18, 20]
    for c_i, w in enumerate(sum_widths, start=1):
        ws_sum.column_dimensions[get_column_letter(c_i)].width = w

    # =========================================================================
    # =========================================================================
    # SHEET 2: Calculation_All_1856
    # =========================================================================
    ws_calc = wb.create_sheet(title=f"{prefix.upper()}_Calculation_All_1856")
    ws_calc.views.sheetView[0].showGridLines = True
    ws_calc.freeze_panes = "A2"

    tp_val = 659.67 if prefix == "ex" else 637.18
    calc_headers = [
        "SL", "Zone Name", "SAP Zone Code", "Zonal Head",
        "Region Name", "SAP Region Code", "Regional Head",
        "SAP Area Code (TTY)", "Territory (Area Name)", "SAP MIO Code", "MIO / Sr. MIO Name", "Designation",
        f"Sep'26 Seen Rx ({brand_title})", "Criteria 1: Seen Rx >= 200?",
        "Aug'26 Total Territory Sales", "Sep'26 Total Territory Sales", "Total Territory Sales Growth",
        "Criteria 2(b): Territory Sales Growth Positive?",
        f"Aug'26 {brand_title} Sales Value", f"Sep'26 {brand_title} Required Sales Value",
        f"Sep'26 {brand_title} Actual Sales Value", f"{brand_title} Value Growth (Sep - Aug)",
        f"{brand_title} TP Price",
        f"Aug'26 {brand_title} Qty (Boxes)", f"Sep'26 {brand_title} Qty (Boxes)",
        f"Required Growth (Box Qty)", f"Actual Growth (Equivalent Box Qty)",
        f"Criteria 2(c): {brand_title} Growth >= 30 Boxes?",
        "All 3 Criteria Met (for 3,000)?",
        f"Eligible {brand_title} Award Tier", f"Award Amount ({brand_title})",
        "RH Recommendation", "RH Cancel Reason",
        "ZH Recommendation", "ZH Cancel Reason",
        "Selected Gift Voucher", "Status", "Timestamp (BDT)"
    ]
    ws_calc.row_dimensions[1].height = 34
    for c_i, h in enumerate(calc_headers, start=1):
        cell = ws_calc.cell(1, c_i, h)
        cell.font = font_hdr
        cell.fill = fill_brand if c_i >= 13 else fill_navy
        cell.alignment = align_cw
        cell.border = border_all

    for idx, r in enumerate(records, start=2):
        ws_calc.row_dimensions[idx].height = 21
        rx_val = r[f"{prefix}_rx"]
        c1_ok = r[f"{prefix}_c1_rx"]
        aug_tot = r["aug_total"]
        sep_tot = r["sep_total"]
        tot_diff = r["tot_growth_val"]
        c2b_ok = r["tot_sales_positive"]
        b_aug = r[f"{prefix}_aug"]
        b_req = r[f"{prefix}_req"]
        b_sep = r[f"{prefix}_sep"]
        b_growth_val = r[f"{prefix}_growth_val"]
        b_aug_box = r.get(f"{prefix}_aug_box")
        b_sep_box = r.get(f"{prefix}_sep_box")
        b_req_box = r.get(f"{prefix}_req_box")
        b_growth_box = r.get(f"{prefix}_growth_box")
        c2c_ok = r[f"{prefix}_growth_met"]
        all3_ok = r[f"{prefix}_all3"]
        tier_str = r[f"{prefix}_tier"]
        aw_amt = r[f"{prefix}_award"]

        row_vals = [
            idx - 1, r["zone_name"], r["sap_zone_code"], r["zm_name"],
            r["region_name"], r["sap_region_code"], r["rm_name"],
            r["tty"], r["area_name"], r["mio_code"], r["mio_name"], r["designation"],
            rx_val, "Yes" if c1_ok else "No",
            aug_tot, sep_tot, tot_diff,
            "Yes" if c2b_ok else ("No (New TTY)" if aug_tot == 0 else "No"),
            b_aug if b_aug is not None else "N/A (New TTY)",
            b_req if b_req is not None else "N/A (New TTY)",
            b_sep,
            b_growth_val if b_growth_val is not None else "N/A",
            tp_val,
            b_aug_box if b_aug_box is not None else "N/A",
            b_sep_box if b_sep_box is not None else "N/A",
            b_req_box if b_req_box is not None else "N/A",
            b_growth_box if b_growth_box is not None else "N/A",
            "Yes" if c2c_ok else ("N/A (New TTY)" if b_req is None else "No"),
            "Yes" if all3_ok else "No",
            tier_str, aw_amt,
            "Recommended" if aw_amt > 0 else "N/A", "",
            "Recommended" if aw_amt > 0 else "N/A", "",
            "", "Pending" if aw_amt > 0 else "N/A", ""
        ]
        for c_i, v in enumerate(row_vals, start=1):
            cell = ws_calc.cell(idx, c_i, v)
            cell.border = border_all
            cell.font = font_data
            if c_i in (2, 4, 5, 7, 9, 11, 12, 30):
                cell.alignment = align_l
            elif c_i in (15, 16, 17, 19, 20, 21, 22):
                cell.alignment = align_r if isinstance(v, (int, float)) else align_c
                if isinstance(v, (int, float)):
                    cell.number_format = "#,##0"
            elif c_i == 23:
                cell.alignment = align_c
                cell.number_format = "0.00"
            else:
                cell.alignment = align_c
                if c_i in (13, 24, 25, 26, 27, 31) and isinstance(v, (int, float)):
                    cell.number_format = "#,##0"

            if c_i in (30, 31):
                if aw_amt == 3000:
                    cell.fill = fill_green
                    cell.font = font_green
                elif aw_amt == 1000:
                    cell.fill = fill_blue_soft
                    cell.font = font_blue
                else:
                    cell.fill = fill_red_soft
                    cell.font = font_red

    calc_widths = [
        6, 18, 13, 24, 20, 14, 24, 14, 22, 13, 25, 24,
        15, 15, 18, 18, 18, 18, 18, 22, 18, 18, 14, 16, 16, 16, 18, 18, 16,
        24, 16, 18, 22, 18, 22, 22, 16, 22
    ]
    for c_i, w in enumerate(calc_widths, start=1):
        ws_calc.column_dimensions[get_column_letter(c_i)].width = w

    # =========================================================================
    # SHEET 3: Eligible Achievers Only (Award_Exium_Sep_2026 / Award_Coralcal_Sep_2026)
    # =========================================================================
    ach_sheet_name = "Award_Exium_Sep_2026" if prefix == "ex" else "Award_Coralcal_Sep_2026"
    ws_ach = wb.create_sheet(title=ach_sheet_name)
    ws_ach.views.sheetView[0].showGridLines = True
    ws_ach.freeze_panes = "A2"

    ach_headers = [
        "SL", "Zone Name", "SAP Zone Code", "Zonal Head",
        "Region Name", "SAP Region Code", "Regional Head",
        "SAP Area Code", "Area Name", "SAP MIO Code", "MIO Name", "Designation",
        "Seen Rx (Sep'26)", "Aug'26 Total Sales", "Sep'26 Total Sales", "Territory Sales Growth", "Total Sales Positive?",
        f"Aug'26 {brand_title} Sales", f"Sep'26 {brand_title} Required", f"Sep'26 {brand_title} Actual", f"{brand_title} Value Growth",
        f"Aug'26 {brand_title} Qty (Boxes)", f"Sep'26 {brand_title} Qty (Boxes)",
        f"Required Growth (Boxes)", f"Actual Growth (Boxes)",
        f"{brand_title} Growth >= 30 Boxes?", "Award Tier", "Award Amount",
        "RH Recommendation", "RH Cancel Reason",
        "ZH Recommendation", "ZH Cancel Reason",
        "Selected Gift Voucher", "Status", "Timestamp (BDT)"
    ]
    ws_ach.row_dimensions[1].height = 32
    for c_i, h in enumerate(ach_headers, start=1):
        cell = ws_ach.cell(1, c_i, h)
        cell.font = font_hdr
        cell.fill = fill_brand
        cell.alignment = align_cw
        cell.border = border_all

    ach_rows = [r for r in records if r[f"{prefix}_award"] > 0]
    for idx, r in enumerate(ach_rows, start=2):
        ws_ach.row_dimensions[idx].height = 21
        aw_amt = r[f"{prefix}_award"]
        row_vals = [
            idx - 1, r["zone_name"], r["sap_zone_code"], r["zm_name"],
            r["region_name"], r["sap_region_code"], r["rm_name"],
            r["tty"], r["area_name"], r["mio_code"], r["mio_name"], r["designation"],
            r[f"{prefix}_rx"], r["aug_total"], r["sep_total"], r["tot_growth_val"], "Yes" if r["tot_sales_positive"] else "No",
            r[f"{prefix}_aug"] if r[f"{prefix}_aug"] is not None else "New TTY",
            r[f"{prefix}_req"] if r[f"{prefix}_req"] is not None else "New TTY",
            r[f"{prefix}_sep"],
            r[f"{prefix}_growth_val"] if r[f"{prefix}_growth_val"] is not None else "New TTY",
            r.get(f"{prefix}_aug_box") if r.get(f"{prefix}_aug_box") is not None else "New TTY",
            r.get(f"{prefix}_sep_box") if r.get(f"{prefix}_sep_box") is not None else "New TTY",
            r.get(f"{prefix}_req_box") if r.get(f"{prefix}_req_box") is not None else "New TTY",
            r.get(f"{prefix}_growth_box") if r.get(f"{prefix}_growth_box") is not None else "New TTY",
            "Yes" if r[f"{prefix}_growth_met"] else "No",
            r[f"{prefix}_tier"], aw_amt,
            "Recommended", "", "Recommended", "", "", "Pending", ""
        ]
        for c_i, v in enumerate(row_vals, start=1):
            cell = ws_ach.cell(idx, c_i, v)
            cell.border = border_all
            cell.font = font_data
            if c_i in (2, 4, 5, 7, 9, 11, 12, 27):
                cell.alignment = align_l
            elif c_i in (14, 15, 16, 18, 19, 20, 21):
                cell.alignment = align_r if isinstance(v, (int, float)) else align_c
                if isinstance(v, (int, float)):
                    cell.number_format = "#,##0"
            else:
                cell.alignment = align_c
                if c_i in (13, 22, 23, 24, 25, 28) and isinstance(v, (int, float)):
                    cell.number_format = "#,##0"

            if c_i in (27, 28):
                cell.fill = fill_green if aw_amt == 3000 else fill_blue_soft
                cell.font = font_green if aw_amt == 3000 else font_blue

    ach_widths = [
        6, 18, 13, 24, 20, 14, 24, 14, 22, 13, 25, 24,
        15, 18, 18, 18, 16, 18, 20, 18, 18, 16, 16, 16, 16, 16, 24, 15, 18, 22, 18, 22, 22, 14, 22
    ]
    for c_i, w in enumerate(ach_widths, start=1):
        ws_ach.column_dimensions[get_column_letter(c_i)].width = w

    wb.save(out_path)
    print(f"Saved: {out_path} ({os.path.getsize(out_path):,} bytes)")


def main():
    with open(DATA_JSON, "r", encoding="utf-8") as f:
        records = json.load(f)

    ex_calc_path = os.path.join(CALC_DIR, "Exium_Award_Calculation_Sep_2026.xlsx")
    cc_calc_path = os.path.join(CALC_DIR, "Coralcal_Award_Calculation_Sep_2026.xlsx")

    build_single_brand_workbook(records, "Exium MUPS", "ex", "1E40AF", "DBEAFE", ex_calc_path)
    build_single_brand_workbook(records, "Coralcal-D/DX/Vita", "cc", "6D28D9", "EDE9FE", cc_calc_path)

    for src, dst_name in [
        (ex_calc_path, "Exium_Award_Calculation_Sep_2026.xlsx"),
        (cc_calc_path, "Coralcal_Award_Calculation_Sep_2026.xlsx"),
    ]:
        dst = os.path.join(BASE_DIR, dst_name)
        try:
            shutil.copy2(src, dst)
            print(f"Copied to {dst}")
        except PermissionError:
            print(f"Skipped copying to {dst} (file currently open in Excel); Calculation/ folder version is updated.")


if __name__ == "__main__":
    main()
