import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
import pandas as pd
import datetime
import os

def generate_september_file():
    folder = r"G:\Exium\2026\Award\September"
    file_june_path = os.path.join(folder, "Exium Award for June'26.xlsx")
    file_ff = os.path.join(folder, "FF list.xlsx")
    output_file = os.path.join(folder, "Exium Award for September'26.xlsx")

    print(f"Reading FF list from: {file_ff}")
    df_ff = pd.read_excel(file_ff)
    print(f"Loaded {len(df_ff)} rows from FF list.")

    print(f"Loading reference June file for exact sheet column widths: {file_june_path}")
    wb_june = openpyxl.load_workbook(file_june_path)

    # Target 35 sheets in exact order
    sheets_order = [
        'B.Baria', 'Barishal-A', 'Barishal-B', 'Bogura', 'Chandpur', 
        'Chattogram-A', 'Chattogram-B', 'Chattogram-C', 'Chattogram-D', 'Cox Bazar', 
        'Cumilla-A', 'Cumilla-B', 'Dhaka-A', 'Dhaka-B', 'Dhaka-C', 
        'Dhaka-D', 'Dhaka-I', 'Dinajpur', 'Faridpur', 'Feni', 
        'Gazipur', 'Jashore', 'Khulna', 'Kushtia', 'Maizdee', 
        'Mymensingh', 'Narayangonj', 'Narsingdi', 'Rajshahi', 'Rangpur', 
        'Savar', 'Sylhet-A', 'Sylhet-B', 'Uttara', 'Tangail'
    ]

    # Mapping dictionary
    def get_ff_zone(sname):
        if sname == 'B.Baria':
            return 'B. Baria (GEN-A)'
        elif sname == 'Cumilla-A':
            return 'Cumilla (GEN-A)'
        else:
            return f"{sname} (GEN-A)"

    def get_col_a_value(sname):
        if sname == 'B.Baria':
            return 'B. Baria'
        else:
            return sname

    # Styles definition
    font_header = Font(name='Calibri', size=11, bold=True)
    font_data = Font(name='Calibri', size=11, bold=False)

    fill_blue = PatternFill(start_color="FFCAEDFB", end_color="FFCAEDFB", fill_type="solid")
    fill_green = PatternFill(start_color="FFC1F0C8", end_color="FFC1F0C8", fill_type="solid")

    align_center = Alignment(horizontal='center', vertical='center')
    align_center_wrap = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center')

    thin_side = Side(style='thin', color='FF000000')
    border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    protect_locked = Protection(locked=True, hidden=False)
    protect_unlocked = Protection(locked=False, hidden=False)

    # September 2026 dates for Row 1
    sept_dates = [datetime.datetime(2026, 9, day, 0, 0) for day in range(1, 31)] # 1 to 30

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active

    total_mios_written = 0

    for s_idx, sname in enumerate(sheets_order):
        ws = wb.create_sheet(title=sname)
        ff_zone_name = get_ff_zone(sname)
        col_a_val = get_col_a_value(sname)

        df_zone = df_ff[df_ff['ZoneName'] == ff_zone_name]
        if len(df_zone) == 0:
            print(f"WARNING: No records found for zone {ff_zone_name} (Sheet: {sname})")

        # --- Set Row 1 (Header) ---
        ws.row_dimensions[1].height = 45.75

        headers = [
            'Zone', 'Zonal Head', 'Region', 'Regional Head',
            'SAP \nArea Code', 'Area Name', 'SAP \nMIO Code', 'MIO/ Sr.MIO Name'
        ]

        # Write A1 to H1
        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_blue
            cell.border = border_all
            cell.protection = protect_locked
            if col_idx in (5, 7):
                cell.alignment = align_center_wrap
            else:
                cell.alignment = align_center

        # Write I1 to AL1 (Dates)
        for d_idx, dt in enumerate(sept_dates, start=9):
            cell = ws.cell(row=1, column=d_idx, value=dt)
            cell.font = font_header
            cell.fill = fill_blue
            cell.border = border_all
            cell.alignment = align_center
            cell.number_format = 'd-mmm'
            cell.protection = protect_locked

        # Write AM1 (Blank)
        cell_am1 = ws.cell(row=1, column=39, value=None)
        cell_am1.font = font_header
        cell_am1.fill = fill_blue
        cell_am1.border = border_all
        cell_am1.alignment = align_center
        cell_am1.number_format = 'd-mmm'
        cell_am1.protection = protect_locked

        # Write AN1 (Total)
        cell_an1 = ws.cell(row=1, column=40, value='Total')
        cell_an1.font = font_header
        cell_an1.fill = fill_green
        cell_an1.border = border_all
        cell_an1.alignment = align_center
        cell_an1.number_format = 'General'
        cell_an1.protection = protect_locked

        # --- Set Data Rows ---
        row_num = 2
        for _, rdata in df_zone.iterrows():
            ws.row_dimensions[row_num].height = 21.75

            # Col A: Zone
            c_a = ws.cell(row=row_num, column=1, value=col_a_val)
            c_a.font = font_data
            c_a.alignment = align_left
            c_a.border = border_all
            c_a.protection = protect_locked

            # Col B: Zonal Head
            c_b = ws.cell(row=row_num, column=2, value=str(rdata['ZMName']).strip())
            c_b.font = font_data
            c_b.alignment = align_left
            c_b.border = border_all
            c_b.protection = protect_locked

            # Col C: Region
            c_c = ws.cell(row=row_num, column=3, value=str(rdata['RegionName']).strip())
            c_c.font = font_data
            c_c.alignment = align_left
            c_c.border = border_all
            c_c.protection = protect_locked

            # Col D: Regional Head
            c_d = ws.cell(row=row_num, column=4, value=str(rdata['RMName']).strip())
            c_d.font = font_data
            c_d.alignment = align_left
            c_d.border = border_all
            c_d.protection = protect_locked

            # Col E: SAP Area Code (Integer)
            c_e = ws.cell(row=row_num, column=5, value=int(rdata['SAPAreaCode']))
            c_e.font = font_data
            c_e.alignment = align_center
            c_e.border = border_all
            c_e.number_format = 'General'
            c_e.protection = protect_locked

            # Col F: Area Name
            c_f = ws.cell(row=row_num, column=6, value=str(rdata['AreaName']).strip())
            c_f.font = font_data
            c_f.alignment = align_left
            c_f.border = border_all
            c_f.protection = protect_locked

            # Col G: SAP MIO Code (Integer)
            c_g = ws.cell(row=row_num, column=7, value=int(rdata['SAPMIOCode']))
            c_g.font = font_data
            c_g.alignment = align_center
            c_g.border = border_all
            c_g.number_format = 'General'
            c_g.protection = protect_locked

            # Col H: MIO/ Sr.MIO Name
            c_h = ws.cell(row=row_num, column=8, value=str(rdata['MIOName']).strip())
            c_h.font = font_data
            c_h.alignment = align_left
            c_h.border = border_all
            c_h.protection = protect_locked

            # Cols I to AL: Daily Rx Input Area (Unlocked!)
            for col_c in range(9, 39): # 9 to 38
                c_in = ws.cell(row=row_num, column=col_c, value=None)
                c_in.font = font_data
                c_in.alignment = align_center
                c_in.border = border_all
                c_in.number_format = '#,##0'
                c_in.protection = protect_unlocked # UNLOCKED FOR USER INPUT!

            # Col AM: Blank Day 31 (Locked!)
            c_am = ws.cell(row=row_num, column=39, value=None)
            c_am.font = font_data
            c_am.alignment = align_center
            c_am.border = border_all
            c_am.number_format = '#,##0'
            c_am.protection = protect_locked # LOCKED

            # Col AN: Total Formula (Locked!)
            c_an = ws.cell(row=row_num, column=40, value=f"=SUM(I{row_num}:AM{row_num})")
            c_an.font = font_data
            c_an.fill = fill_green
            c_an.alignment = align_center
            c_an.border = border_all
            c_an.number_format = '#,##0'
            c_an.protection = protect_locked # LOCKED

            row_num += 1
            total_mios_written += 1

        # --- Apply Exact Column Widths and Hidden States from June File ---
        ws_june = wb_june[sname] if sname in wb_june.sheetnames else None

        # Helper to get June width with sensible fallback
        def get_june_width(col_letter, fallback):
            if ws_june and col_letter in ws_june.column_dimensions:
                w = ws_june.column_dimensions[col_letter].width
                if w is not None:
                    return w
            return fallback

        # Columns A to H widths
        ws.column_dimensions['A'].width = get_june_width('A', 11.0)
        
        # Column B (Zonal Head) - HIDDEN like June file
        ws.column_dimensions['B'].width = get_june_width('B', 17.14)
        ws.column_dimensions['B'].hidden = True

        ws.column_dimensions['C'].width = get_june_width('C', 18.0)
        ws.column_dimensions['D'].width = get_june_width('D', 23.0)
        ws.column_dimensions['E'].width = get_june_width('E', 9.71)
        ws.column_dimensions['F'].width = get_june_width('F', 18.0)
        ws.column_dimensions['G'].width = get_june_width('G', 10.29)
        ws.column_dimensions['H'].width = get_june_width('H', 28.0)

        # Columns I to AL (Cols 9 to 38) - Date columns: width 8.0
        # (Guarantees '1-Sep' to '30-Sep' is completely visible with no '###' overflow, while keeping columns compact)
        for c in range(9, 39):
            col_letter = get_column_letter(c)
            ws.column_dimensions[col_letter].width = 8.0

        # Column AM (Col 39) - Blank 31st Day: width 13.0 and HIDDEN like June file
        ws.column_dimensions['AM'].width = get_june_width('AM', 13.0)
        ws.column_dimensions['AM'].hidden = True

        # Column AN (Col 40) - Total: width 8.71 like June file
        ws.column_dimensions['AN'].width = get_june_width('AN', 8.71)

        # Freeze Panes at I2
        ws.freeze_panes = 'I2'

        # Ensure Grid lines are visible
        ws.views.sheetView[0].showGridLines = True

        # Sheet Protection Configuration:
        # We leave ws.protection.sheet = False in the Excel file so that Google Sheets
        # does not globally lock all cells upon upload.
        ws.protection.sheet = False

    # Remove the initial default sheet
    wb.remove(default_sheet)

    print(f"Total MIO records written across all 35 sheets: {total_mios_written}")
    print(f"Saving workbook to: {output_file}")
    wb.save(output_file)
    print("Successfully generated and saved Exium Award for September'26.xlsx!")

if __name__ == '__main__':
    generate_september_file()
